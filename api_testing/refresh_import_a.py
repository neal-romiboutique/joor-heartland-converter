#!/usr/bin/env python3
"""
Refresh the Import A batch files against live Heartland before importing.

The batch files are a ~2-week-old snapshot. Every pass-through column in them
(cost, price, vendor, dimensions, active) would be written back to Heartland
verbatim, reverting anything changed since the export. This rewrites those
columns from current API values, so the files carry only the intended change:
Description + Style Name to title case.

Casing is NOT recomputed. If the live text matches the batch file's text
ignoring case/whitespace, the batch file's title-cased value is reused. If it
differs, the title was edited since export and the row is held back.

READ-ONLY against Heartland. Writes new local files; originals untouched.

Usage:
    python3 refresh_import_a.py /path/to/batch_folder
"""
import json, os, re, sys, time
from pathlib import Path

import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# ----------------------------- config ---------------------------------------
SLEEP = 0.05           # polite pacing between requests
MAX_RETRIES = 4
CACHE_EVERY = 200      # checkpoint to disk every N fetched grids
TEXT_COLS = ("Item Grid #", "Description", "Style Name",
             "Dimension 1 Values", "Dimension 2 Values",
             "Primary Vendor", "Item Primary Vendor")

# ----------------------------- credentials ----------------------------------
def load_env():
    p = Path(__file__).resolve().parent / ".env"
    if p.exists():
        for line in p.read_text().splitlines():
            if line.strip() and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip())

load_env()
TOKEN = os.environ.get("HEARTLAND_TOKEN")
SUB = os.environ.get("HEARTLAND_SUBDOMAIN")
if not (TOKEN and SUB):
    sys.exit("ERROR: HEARTLAND_TOKEN / HEARTLAND_SUBDOMAIN missing from .env")
BASE = f"https://{SUB}.retail.heartland.us/api"
HEADERS = {"Authorization": f"Bearer {TOKEN}"}

try:
    from vendor_id_map import VENDOR_ID_MAP
except ImportError:
    sys.exit("ERROR: vendor_id_map.py not found next to this script.")

# id -> name, for the case where a grid's vendor changed since export.
ID_TO_VENDOR, _dupe_ids = {}, set()
for _n, _i in VENDOR_ID_MAP.items():
    if _i in ID_TO_VENDOR and ID_TO_VENDOR[_i] != _n:
        _dupe_ids.add(_i)
    ID_TO_VENDOR.setdefault(_i, _n)

# ----------------------------- helpers --------------------------------------
def norm(s) -> str:
    """Case- and whitespace-insensitive comparison key."""
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()

def grid_key(v) -> str:
    """'738', 738, 738.0 -> '738'."""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s

def fetch_grid(gid, cache):
    if gid in cache:
        return cache[gid]
    for attempt in range(MAX_RETRIES):
        try:
            r = requests.get(f"{BASE}/item_grids/{gid}",
                             headers=HEADERS, timeout=30)
        except requests.RequestException as e:
            if attempt == MAX_RETRIES - 1:
                return {"_error": f"network: {e}"}
            time.sleep(2 ** attempt)
            continue
        if r.status_code == 200:
            d = r.json()
            d = d.get("result", d)
            cache[gid] = d
            return d
        if r.status_code == 404:
            cache[gid] = {"_error": "not found"}
            return cache[gid]
        if r.status_code in (429, 500, 502, 503, 504):
            time.sleep(2 ** attempt)
            continue
        return {"_error": f"HTTP {r.status_code}"}
    return {"_error": "retries exhausted"}

def build_row(old, grid, headers):
    """Return (new_row_dict, holdback_reason_or_None, [drift strings])."""
    if "_error" in grid:
        return None, grid["_error"], []
    if not grid.get("active?", True):
        return None, "inactive in Heartland", []

    dims = grid.get("dimensions") or []
    if len(dims) < 2:
        return None, f"expected 2 dimensions, found {len(dims)}", []
    for d in dims[:2]:
        for v in (d.get("values") or []):
            if "," in str(v):
                return None, f"comma inside dimension value '{v}'", []

    # --- Description / Style Name: reuse batch casing only if text unchanged
    live_desc = grid.get("description") or ""
    if norm(live_desc) != norm(old.get("Description")):
        return None, (f"description text changed since export "
                      f"(live: {live_desc!r})"), []
    live_style = (grid.get("item_custom") or {}).get("style") or ""
    if norm(live_style) != norm(old.get("Style Name")):
        return None, (f"style name text changed since export "
                      f"(live: {live_style!r})"), []

    # --- vendor: keep the batch string verbatim if the id still agrees
    vid = grid.get("item_primary_vendor_id")
    old_vendor = str(old.get("Primary Vendor") or "").strip()
    if VENDOR_ID_MAP.get(old_vendor) == vid:
        vendor = old_vendor
    elif vid in ID_TO_VENDOR and vid not in _dupe_ids:
        vendor = ID_TO_VENDOR[vid]
    else:
        return None, f"vendor id {vid} unmapped or ambiguous", []

    new = {
        "Item Grid #":         grid_key(grid.get("public_id")),
        "Description":         old.get("Description"),
        "Style Name":          old.get("Style Name"),
        "Default Cost":        grid.get("item_cost"),
        "Price":               grid.get("item_price"),
        "Primary Vendor":      vendor,
        "Item Primary Vendor": vendor,
        "Dimension 1":         dims[0].get("name"),
        "Dimension 1 Values":  ",".join(str(v) for v in dims[0].get("values") or []),
        "Dimension 2":         dims[1].get("name"),
        "Dimension 2 Values":  ",".join(str(v) for v in dims[1].get("values") or []),
        "Active":              "true",
        "Generate Items":      "true",
    }

    drift = []
    for h in headers:
        if h in ("Description", "Style Name", "Generate Items"):
            continue
        o, n = old.get(h), new.get(h)
        if norm(o) != norm(n):
            drift.append(f"{h}: {o!r} -> {n!r}")
    return new, None, drift

def write_xlsx(path, headers, rows):
    wb = Workbook(); ws = wb.active
    bold = Font(name="Arial", size=10, bold=True)
    norm_f = Font(name="Arial", size=10)
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = bold
    for r_i, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=r_i, column=c)
            v = row.get(h)
            if h in TEXT_COLS and v is not None:
                cell.value = str(v); cell.number_format = "@"
            else:
                cell.value = v
            cell.font = norm_f
    for c, h in enumerate(headers, 1):
        vals = [h] + [str(r.get(h) or "") for r in rows]
        ws.column_dimensions[get_column_letter(c)].width = \
            min(max(len(s) for s in vals) + 2, 34)
    ws.freeze_panes = "A2"
    wb.save(path)

# ----------------------------- main -----------------------------------------
def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 refresh_import_a.py /path/to/batch_folder")
    folder = Path(sys.argv[1]).expanduser().resolve()
    batches = sorted(folder.glob("IMPORT_A_batch_*_of_9.xlsx"))
    batches = [b for b in batches if "REFRESHED" not in b.name]
    if not batches:
        sys.exit(f"No IMPORT_A_batch_*_of_9.xlsx files found in {folder}")
    print(f"Found {len(batches)} batch file(s) in {folder}")

    cache_path = folder / ".refresh_cache.json"
    cache = json.loads(cache_path.read_text()) if cache_path.exists() else {}
    if cache:
        print(f"Resuming with {len(cache)} grid(s) already cached.")

    held, fetched = [], 0
    for bf in batches:
        wb = load_workbook(bf, data_only=True, read_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        old_rows = [dict(zip(headers, r)) for r in it
                    if any(c is not None and str(c).strip() for c in r)]
        wb.close()

        new_rows = []
        for old in old_rows:
            gid = grid_key(old.get("Item Grid #"))
            if not gid:
                continue
            before = gid in cache
            grid = fetch_grid(gid, cache)
            if not before:
                fetched += 1
                time.sleep(SLEEP)
                if fetched % CACHE_EVERY == 0:
                    cache_path.write_text(json.dumps(cache))
                    print(f"  ... {fetched} fetched")
            row, reason, drift = build_row(old, grid, headers)
            if reason:
                held.append({"Batch": bf.name, "Item Grid #": gid,
                             "Reason": reason, "Drift": ""})
            else:
                new_rows.append(row)
                if drift:
                    held.append({"Batch": bf.name, "Item Grid #": gid,
                                 "Reason": "REFRESHED (would have reverted)",
                                 "Drift": "; ".join(drift)})

        out = bf.with_name(bf.stem + "_REFRESHED.xlsx")
        write_xlsx(out, headers, new_rows)
        cache_path.write_text(json.dumps(cache))
        print(f"{bf.name}: {len(old_rows)} in -> {len(new_rows)} out "
              f"({len(old_rows) - len(new_rows)} held back)  ->  {out.name}")

    if held:
        log = folder / "REFRESH_LOG.xlsx"
        write_xlsx(log, ["Batch", "Item Grid #", "Reason", "Drift"], held)
        reverts = sum(1 for h in held if h["Reason"].startswith("REFRESHED"))
        drops = len(held) - reverts
        print(f"\n{reverts} row(s) had stale values corrected.")
        print(f"{drops} row(s) held back for review.")
        print(f"See {log.name}")
    else:
        print("\nNo drift and no held-back rows.")
    print("\nImport the *_REFRESHED.xlsx files. Do not open them in Excel.")

if __name__ == "__main__":
    main()
